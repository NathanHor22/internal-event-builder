import csv
import io
import os
import json
from datetime import datetime
from database import get_db
import config

def get_event_content_data(event_id):
    db = get_db()
    event = db.execute("SELECT * FROM events WHERE id = ?", (event_id,)).fetchone()
    if not event:
        db.close()
        return None, None, None
    campaigns = db.execute("SELECT * FROM campaigns WHERE event_id = ? ORDER BY created_at", (event_id,)).fetchall()
    pieces = db.execute("""
        SELECT cp.*, c.name as campaign_name, c.phase as campaign_phase
        FROM content_pieces cp
        JOIN campaigns c ON cp.campaign_id = c.id
        WHERE c.event_id = ?
        ORDER BY c.id, cp.sort_order, cp.created_at
    """, (event_id,)).fetchall()
    platforms = db.execute("SELECT * FROM event_platforms WHERE event_id = ? AND is_selected = 1", (event_id,)).fetchall()
    db.close()
    return dict(event), [dict(c) for c in campaigns], [dict(p) for p in pieces], [dict(pl) for pl in platforms]

def export_csv(event_id):
    result = get_event_content_data(event_id)
    if not result[0]:
        return None
    event, campaigns, pieces, platforms = result
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(['Campaign', 'Phase', 'Title', 'Format', 'Platform', 'Copywriting',
                     'Video Script', 'Caption', 'Timing', 'Is Ad', 'Ad Platform',
                     'Ad Budget', 'Ad Start', 'Ad End', 'Status'])
    for p in pieces:
        writer.writerow([
            p.get('campaign_name', ''), p.get('campaign_phase', ''),
            p.get('title', ''), p.get('format', ''), p.get('platform', ''),
            p.get('copywriting', ''), p.get('video_script', ''), p.get('caption', ''),
            p.get('timing', ''), 'Yes' if p.get('is_ad') else 'No',
            p.get('ad_platform', ''), p.get('ad_budget', ''),
            p.get('ad_start_date', ''), p.get('ad_end_date', ''), p.get('status', '')
        ])
    return output.getvalue()

def export_excel(event_id):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    result = get_event_content_data(event_id)
    if not result[0]:
        return None
    event, campaigns, pieces, platforms = result

    wb = Workbook()
    # Summary sheet
    ws_summary = wb.active
    ws_summary.title = "Event Overview"
    ws_summary.column_dimensions['A'].width = 20
    ws_summary.column_dimensions['B'].width = 60

    header_font = Font(bold=True, size=14)
    label_font = Font(bold=True)

    ws_summary['A1'] = event.get('name', 'Untitled Event')
    ws_summary['A1'].font = header_font
    row = 3
    for label, key in [('Description', 'description'), ('Target Audience', 'target_audience'),
                        ('Theme', 'theme'), ('Location', 'location'), ('Status', 'status')]:
        ws_summary.cell(row=row, column=1, value=label).font = label_font
        ws_summary.cell(row=row, column=2, value=event.get(key, ''))
        row += 1

    # Content sheet
    ws = wb.create_sheet("Content Plan")
    headers = ['Campaign', 'Phase', 'Title', 'Format', 'Platform', 'Copywriting',
               'Video Script', 'Caption', 'Timing', 'Is Ad', 'Ad Platform',
               'Ad Budget', 'Ad Start', 'Ad End', 'Status']

    header_fill = PatternFill(start_color='6366F1', end_color='6366F1', fill_type='solid')
    header_font_white = Font(bold=True, color='FFFFFF')

    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = header_font_white
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center')

    for i, p in enumerate(pieces, 2):
        ws.cell(row=i, column=1, value=p.get('campaign_name', ''))
        ws.cell(row=i, column=2, value=p.get('campaign_phase', ''))
        ws.cell(row=i, column=3, value=p.get('title', ''))
        ws.cell(row=i, column=4, value=p.get('format', ''))
        ws.cell(row=i, column=5, value=p.get('platform', ''))
        ws.cell(row=i, column=6, value=p.get('copywriting', ''))
        ws.cell(row=i, column=7, value=p.get('video_script', ''))
        ws.cell(row=i, column=8, value=p.get('caption', ''))
        ws.cell(row=i, column=9, value=p.get('timing', ''))
        ws.cell(row=i, column=10, value='Yes' if p.get('is_ad') else 'No')
        ws.cell(row=i, column=11, value=p.get('ad_platform', ''))
        ws.cell(row=i, column=12, value=p.get('ad_budget', ''))
        ws.cell(row=i, column=13, value=p.get('ad_start_date', ''))
        ws.cell(row=i, column=14, value=p.get('ad_end_date', ''))
        ws.cell(row=i, column=15, value=p.get('status', ''))

    for col in range(1, 16):
        ws.column_dimensions[ws.cell(row=1, column=col).column_letter].width = 18

    filepath = os.path.join(config.EXPORT_DIR, f"event_{event_id}_{datetime.now().strftime('%Y%m%d%H%M%S')}.xlsx")
    wb.save(filepath)
    return filepath

def export_pdf(event_id):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer

    result = get_event_content_data(event_id)
    if not result[0]:
        return None
    event, campaigns, pieces, platforms = result

    filepath = os.path.join(config.EXPORT_DIR, f"event_{event_id}_{datetime.now().strftime('%Y%m%d%H%M%S')}.pdf")
    doc = SimpleDocTemplate(filepath, pagesize=landscape(A4), topMargin=0.5*inch, bottomMargin=0.5*inch)
    styles = getSampleStyleSheet()
    elements = []

    # Title
    title_style = ParagraphStyle('CustomTitle', parent=styles['Title'], fontSize=20, spaceAfter=12)
    elements.append(Paragraph(event.get('name', 'Content Plan'), title_style))
    elements.append(Spacer(1, 12))

    # Event info
    info_style = ParagraphStyle('Info', parent=styles['Normal'], fontSize=10, spaceAfter=4)
    for label, key in [('Target Audience', 'target_audience'), ('Theme', 'theme'),
                        ('Location', 'location'), ('Status', 'status')]:
        val = event.get(key, '')
        if val:
            elements.append(Paragraph(f"<b>{label}:</b> {val}", info_style))

    if platforms:
        plat_str = ", ".join(p.get('platform', '') for p in platforms)
        elements.append(Paragraph(f"<b>Platforms:</b> {plat_str}", info_style))

    elements.append(Spacer(1, 20))

    # Content table
    if pieces:
        cell_style = ParagraphStyle('Cell', parent=styles['Normal'], fontSize=7, leading=9)
        table_data = [['Campaign', 'Title', 'Format', 'Platform', 'Copy', 'Caption', 'Ad?', 'Budget', 'Status']]
        for p in pieces:
            table_data.append([
                Paragraph(str(p.get('campaign_name', '')), cell_style),
                Paragraph(str(p.get('title', '')), cell_style),
                str(p.get('format', '')),
                str(p.get('platform', '')),
                Paragraph(str(p.get('copywriting', ''))[:200], cell_style),
                Paragraph(str(p.get('caption', ''))[:150], cell_style),
                'Yes' if p.get('is_ad') else 'No',
                f"${p['ad_budget']:.0f}" if p.get('ad_budget') else '',
                str(p.get('status', ''))
            ])

        col_widths = [1.2*inch, 1.5*inch, 0.7*inch, 0.8*inch, 2.5*inch, 2*inch, 0.5*inch, 0.7*inch, 0.7*inch]
        t = Table(table_data, colWidths=col_widths, repeatRows=1)
        t.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#6366F1')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTSIZE', (0, 0), (-1, 0), 8),
            ('FONTSIZE', (0, 1), (-1, -1), 7),
            ('ALIGN', (0, 0), (-1, 0), 'CENTER'),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CCCCCC')),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F5F5FF')]),
            ('VALIGN', (0, 0), (-1, -1), 'TOP'),
            ('TOPPADDING', (0, 0), (-1, -1), 4),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ]))
        elements.append(t)

    doc.build(elements)
    return filepath
