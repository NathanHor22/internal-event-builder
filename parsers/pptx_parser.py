def parse_pptx(filepath):
    try:
        from pptx import Presentation
        prs = Presentation(filepath)
        text_parts = []
        for i, slide in enumerate(prs.slides, 1):
            slide_texts = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for para in shape.text_frame.paragraphs:
                        text = para.text.strip()
                        if text:
                            slide_texts.append(text)
                if shape.has_table:
                    for row in shape.table.rows:
                        row_text = " | ".join(cell.text.strip() for cell in row.cells)
                        if row_text.strip(" |"):
                            slide_texts.append(row_text)
            if slide_texts:
                text_parts.append(f"--- Slide {i} ---\n" + "\n".join(slide_texts))
            # Notes
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                notes = slide.notes_slide.notes_text_frame.text.strip()
                if notes:
                    text_parts.append(f"[Speaker Notes] {notes}")
        return "\n\n".join(text_parts)
    except Exception:
        return ""
